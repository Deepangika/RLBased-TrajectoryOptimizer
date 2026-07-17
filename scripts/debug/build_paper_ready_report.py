from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "outputs" / "full_realisability_matrix"
FIG = DATA / "figures"
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

summary = list(csv.DictReader((DATA / "realisability_summary.csv").open(encoding="utf-8")))
matrix_rows = list(csv.DictReader((DATA / "full_realisability_matrix.csv").open(encoding="utf-8")))

GESTURES = ["wave", "reach", "point"]
STATES = ["confident", "calm", "hesitant", "friendly", "confused", "angry"]
FEATURES = ["weight", "time", "flow_boundness", "space_indirectness", "shape_arcness"]
TARGETS = {
    state: {k: float(next(r for r in matrix_rows if r["state"] == state)[f"target_{k}"]) for k in FEATURES}
    for state in STATES
}

NAVY = RGBColor(20, 48, 78)
BLUE = RGBColor(38, 92, 140)
MUTED = RGBColor(93, 106, 120)
LIGHT = "EAF1F7"
LIGHT2 = "F4F6F8"
WHITE = "FFFFFF"


def set_font(run, name="Aptos", size=10.2, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[ci])
            shade(cell, LIGHT if ri == 0 else WHITE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run, size=8.6, bold=(ri == 0), color=NAVY if ri == 0 else None)
    set_repeat_table_header(table.rows[0])


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(text), size=9, italic=True, color=MUTED)
    return p


def add_figure(doc, filename, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / filename), width=Inches(width))
    add_caption(doc, caption)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True, color=NAVY)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(text))
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(text), name="Cambria Math", size=11, italic=True, color=NAVY)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.86)
section.right_margin = Inches(0.86)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.2)
normal.paragraph_format.space_after = Pt(5.5)
normal.paragraph_format.line_spacing = 1.08
for name, size, color, before, after in [
    ("Title", 24, NAVY, 0, 8),
    ("Subtitle", 12, MUTED, 0, 18),
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 12.5, NAVY, 10, 5),
    ("Heading 3", 11, NAVY, 8, 4),
]:
    style = styles[name]
    style.font.name = "Aptos Display" if name in ["Title", "Heading 1"] else "Aptos"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = name not in ["Subtitle"]
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("EXPRESSIVE ROBOT GESTURE GENERATION  |  TECHNICAL SECTION"), size=8.2, bold=True, color=MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Spatiotemporal Laban-feature optimisation and realisability evaluation"), size=8, color=MUTED)

title = doc.add_paragraph(style="Title")
set_font(title.add_run("Spatiotemporal Laban-Feature Optimisation"), name="Aptos Display", size=24, bold=True, color=NAVY)
subtitle = doc.add_paragraph(style="Subtitle")
set_font(subtitle.add_run("Algorithm, target-realisability protocol, and empirical results across gestures and affective states"), size=12, color=MUTED)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
set_font(p.add_run("Paper-ready draft section  •  90 optimisation runs  •  3 gestures  •  6 states  •  5 random seeds"), size=9.5, bold=True, color=BLUE)

doc.add_heading("Overview", level=1)
add_body(doc, "This section describes a direct spatiotemporal optimiser that modifies a reference robot-arm gesture to approximate a target Laban Movement Analysis profile while retaining the gesture’s functional structure. The method separates geometric styling from temporal retiming, evaluates each candidate using normalised trajectory-derived Laban proxies, and balances target matching against path preservation, smoothness, and joint feasibility. A multi-seed realisability matrix was used to distinguish profiles that were consistently achieved from those affected by feature coupling, normalisation saturation, or stochastic search failure.")

doc.add_heading("Optimisation method", level=1)
doc.add_heading("Trajectory representation", level=2)
add_body(doc, "Let q_ref(u) ∈ ℝ² denote the reference shoulder–elbow trajectory, where u ∈ [0,1] is normalised gesture progress. Rather than optimising every joint value at every frame, the spatial residual is represented using K_s smooth sine basis functions for each joint:")
add_equation(doc, "q_spatial,j(u) = q_ref,j(u) + Δ_max Σₖ c_jk sin(kπu) + r(u)e_jΔ_end")
add_body(doc, "Here, c_jk ∈ [−1,1] is a spatial coefficient, Δ_max scales each basis contribution into radians, e_j is an optional final-pose offset, and r(u)=u²(3−2u) introduces that offset smoothly. With six spatial bases, two joints, two endpoint offsets, and four timing bases, the optimiser searches an 18-dimensional coefficient vector. The sine bases preserve the initial and final poses by construction; soft endpoint offsets deliberately relax only the final-pose constraint.")

doc.add_heading("Monotonic temporal retiming", level=2)
add_body(doc, "A shared positive speed profile is constructed from K_t timing coefficients. Exponentiation guarantees positive progress, after which cumulative normalisation maps the result back to [0,1]:")
add_equation(doc, "s(u)=exp(α Σₖ dₖ sin(kπu)),     τ(u)=cumsum(s(u))/cumsum(s(1))")
add_body(doc, "The final variant is q_var(u)=q_spatial(τ(u)). This changes the distribution of speed within the gesture while preserving total duration. Consequently, the current implementation can express local acceleration, deceleration, and hesitation, but cannot make the complete gesture globally shorter or longer.")

doc.add_heading("Objective function", level=2)
add_body(doc, "For each candidate, raw movement features are calculated and mapped through gesture-specific empirical min–max ranges. The primary target-matching term is the weighted root-mean-square error between the candidate’s normalised feature vector y(q_var) and the target profile y*:")
add_equation(doc, "L_feature = √[Σᵢ wᵢ(yᵢ−yᵢ*)² / Σᵢwᵢ]")
add_body(doc, "The total objective augments this term with penalties for joint-space deviation, nearest-path distance, endpoint and direction error, path-length change, excessive detours, maximum path deviation, jerk, joint-limit violations, coefficient magnitude, and temporal-warp roughness:")
add_equation(doc, "L_total = L_feature + Σₘ λₘLₘ")
add_body(doc, "Differential Evolution performs the global search within coefficient bounds [−1,1], followed by bounded L-BFGS-B polishing. The final reward stored by the wider pipeline is the negative loss; however, the procedure itself is direct trajectory optimisation rather than reinforcement learning.")

doc.add_heading("Algorithm", level=2)
algorithm = [
    "Load the reference gesture, empirical normalisation ranges, and target Laban profile.",
    "Construct spatial and timing sine bases and initialise a bounded Differential Evolution population.",
    "For each coefficient vector, add spatial residuals, apply the optional endpoint offset, and generate a monotonic time warp.",
    "Compute the modified trajectory’s raw and normalised Laban features.",
    "Evaluate feature error, gesture-preservation penalties, smoothness, and joint feasibility.",
    "Update the population, repeat for the specified generations, and polish the best candidate locally.",
    "Save the best trajectory, coefficients, raw and normalised features, and diagnostic metrics.",
]
for i, step in enumerate(algorithm, 1):
    p = doc.add_paragraph(style="List Number")
    set_font(p.add_run(step))

doc.add_heading("Target-realisability evaluation", level=1)
doc.add_heading("Experimental design", level=2)
add_body(doc, "The evaluation tested whether the optimiser could realise six predefined affective profiles for three reference gestures (wave, reach, and point). Five independent random seeds were used for every gesture–state combination, yielding 3×6×5=90 optimisation runs. Each run used 45 Differential Evolution generations, a population multiplier of five, and up to 100 local-polishing iterations.")

table = doc.add_table(rows=1, cols=6)
headers = ["State", "Weight", "Time", "Flow", "Space", "Shape"]
for i, h in enumerate(headers): table.rows[0].cells[i].text = h
for state in STATES:
    cells = table.add_row().cells
    vals = TARGETS[state]
    content = [state.title()] + [f"{vals[k]:.2f}" for k in FEATURES]
    for i, value in enumerate(content): cells[i].text = value
style_table(table, widths=[1.35, 1.0, 1.0, 1.0, 1.0, 1.0])
add_caption(doc, "Table 1. Predefined normalised target profiles evaluated in the realisability matrix.")

doc.add_heading("Validity and success criteria", level=2)
add_body(doc, "A key extension to the original code recorded both clipped and unclipped normalised features. Clipped values remain useful to keep optimisation bounded, but unclipped values reveal whether an apparent boundary value (e.g., 1.0) actually lies outside the empirical range. A run was classified as realised only when all five unclipped feature errors were finite and no greater than 0.10. Path preservation was recorded separately using a path-length ratio interval of [0.70,1.30], and joint feasibility required zero joint-limit penalty. These thresholds are engineering diagnostics rather than perceptually validated equivalence margins.")

doc.add_heading("Results", level=1)
add_figure(doc, "fig1_success_rate_heatmap.png", "Figure 1. Proportion of five random-seed runs in which all unclipped feature errors were ≤0.10.", width=5.75)
add_body(doc, "Across all 90 runs, 37 (41.1%) met the five-feature realisation criterion and 34 (37.8%) additionally satisfied the path-length and joint-limit checks. Reliability varied strongly by gesture and state. Wave achieved confidence, calm, friendliness, and confusion in every seed, whereas hesitation and anger were not realised. Reach was reliable for confidence and friendliness, moderately reliable for calm and confusion, and rarely successful for hesitation. Point was substantially less stable: only one seed succeeded for confidence, friendliness, and confusion, while calm, hesitation, and anger were never realised.")

table = doc.add_table(rows=1, cols=7)
for i, h in enumerate(["Gesture", *[s.title() for s in STATES]]): table.rows[0].cells[i].text = h
for gesture in GESTURES:
    cells = table.add_row().cells
    cells[0].text = gesture.title()
    for si, state in enumerate(STATES, 1):
        r = next(x for x in summary if x["gesture"] == gesture and x["state"] == state)
        cells[si].text = f"{int(r['realised_runs'])}/5"
style_table(table, widths=[1.05, 0.87, 0.75, 0.9, 0.85, 0.9, 0.75])
add_caption(doc, "Table 2. Successful target realisations across five seeds for each gesture–state combination.")

add_figure(doc, "fig2_median_rmse_heatmap.png", "Figure 2. Median RMSE computed using unclipped normalised features. Lower values indicate closer target matching.")
add_body(doc, "Successful wave combinations typically produced near-zero median RMSE, demonstrating that the 18-dimensional basis representation can express those profiles without consistently approaching coefficient bounds. Reach displayed a bimodal pattern: several seeds converged closely while others entered saturated regions. Point showed the strongest seed dependence and the largest median errors, indicating that its objective landscape contains difficult basins rather than proving that all failed targets are structurally impossible.")

add_figure(doc, "fig3_clipping_rate_heatmap.png", "Figure 3. Fraction of runs containing at least one feature outside its empirical normalisation range.")
add_body(doc, "Forty-two runs (46.7%) contained at least one clipped feature. Saturation was absent in all wave runs except anger, but frequent for point and for failed reach runs. Because clipping maps all values beyond an empirical boundary to exactly zero or one, it creates flat objective regions in which physically different extreme trajectories become numerically indistinguishable. This explains why removing preservation penalties degraded rather than improved earlier feature-only tests: the preservation terms also regularise the search and discourage entry into saturated regions.")

add_figure(doc, "fig4_feature_error_heatmap.png", "Figure 4. Mean absolute target error by feature and gesture, calculated before clipping.")
add_body(doc, "For wave, the principal remaining mismatch was target-dependent: hesitation consistently missed Flow and, to a lesser extent, Time, while angry runs produced extreme Weight and Time values outside the calibrated range. For point, errors were broad across Weight, Time, Flow, and Space, whereas Shape was often matched closely. This pattern suggests that Shape is comparatively controllable under the current basis, while the dynamic features and global Space proxy interact strongly with temporal warping and empirical clipping.")

add_figure(doc, "fig5_best_profiles.png", "Figure 5. Best solution found for each gesture–state combination compared with the requested five-dimensional profile.", width=6.15)
add_body(doc, "The best-run comparison is essential for separating representational feasibility from search reliability. Confidence was matched by all three gestures at least once, showing that the profile lies within the current representation’s achievable set. Friendly and confused profiles were also matched by point in one seed despite poor median performance. These combinations should therefore be described as achievable but search-inconsistent. In contrast, no angry run met the criterion for any gesture, and even the best wave and reach solutions retained substantial error. Anger is consequently unresolved under the present experiment: it may be representation-limited, conflict with preservation penalties, or remain inaccessible because the clipped objective impairs search.")

doc.add_heading("Interpretation and limitations", level=1)
add_body(doc, "Target realisability should not be inferred from optimiser termination alone. The results support a three-way distinction:")
add_bullet(doc, "Reliably realised: high success across seeds with low unclipped error (e.g., wave–confident and reach–confident).")
add_bullet(doc, "Achievable but search-inconsistent: at least one close solution exists, but success is seed-sensitive (e.g., point–confident, point–friendly, and point–confused).")
add_bullet(doc, "Not realised under the tested configuration: no seed meets tolerance; additional relaxed-but-regularised tests are needed before claiming structural infeasibility (notably angry for all gestures and hesitant for wave/point).")

add_body(doc, "Several limitations qualify these findings. First, the normalisation ranges are empirical and gesture-specific; identical normalised values need not correspond to identical raw movement intensity across gestures. Second, the target profiles are prescribed rather than perceptually calibrated, so numerical matching does not guarantee that humans recognise the intended state. Third, the arm is a planar two-link model without torque, collision, velocity, or full-embodiment constraints. Fourth, the current time warp redistributes motion within a fixed duration and cannot express overall duration changes. Fifth, Space indirectness is based on a path-length-to-displacement proxy that is unstable for cyclic gestures and should be replaced or complemented by a bounded local path-efficiency measure.")

doc.add_heading("Recommended reporting statement", level=1)
add_body(doc, "The optimiser demonstrated that a compact spatiotemporal basis can realise several predefined Laban profiles while approximately preserving reference-gesture geometry. Across 90 runs, success depended on both gesture and target state. Wave profiles were generally the most reliable, reach profiles were moderately reliable with sufficient search budget, and point profiles exhibited pronounced seed sensitivity. Recording unclipped normalised features revealed that almost half of the runs entered out-of-range regions hidden by conventional [0,1] clipping. Consequently, failed runs cannot be interpreted as direct evidence that a target is physically unrealizable. Instead, the matrix identifies combinations that are reliably achieved, achievable but search-inconsistent, or unresolved under the current representation and objective. Human and VLM evaluation remains necessary to determine whether numerically matched profiles produce perceptually legible affective gestures.")

doc.add_heading("Reproducibility details", level=1)
add_body(doc, "Reference trajectories contained 160 samples over 2 s for a planar arm with link lengths 0.30 m and 0.25 m. The optimiser used six spatial bases per joint, four temporal bases, a maximum per-basis spatial scale of 0.35 rad, a maximum endpoint-offset scale of 0.22 rad, and soft endpoint handling. The matrix used seeds {7,17,27,37,47}. Per-run raw data, unclipped and clipped features, preservation measures, and summary statistics were retained to enable threshold sensitivity analysis and future comparison with perceptual judgements.")

path = OUT / "Paper_Ready_Optimiser_Method_and_Realisability_Results.docx"
doc.save(path)
print(path)
